"""API·권한 테스트 (CLAUDE.md 8.3절, 10장): 역할별 허용/거부, IDOR, 공유 철회."""
import os

import pytest

from tests.conftest import H

from app.config import Settings, validate_settings
from app.seed import seed


def _demo_id(c):
    r = c.get("/api/projects?scope=mine")
    assert r.status_code == 200
    return next(p["id"] for p in r.json() if p["name"] == "Poly-Si 식각 레시피 개발 (데모)")


def test_requires_login(seeded):
    from fastapi.testclient import TestClient
    from app.main import app
    with TestClient(app) as c:
        assert c.get("/api/projects").status_code == 401


def test_csrf_header_required(client_for):
    c = client_for("E1001")
    r = c.post("/api/projects", json={}, headers={"X-Requested-With": ""})
    assert r.status_code == 403


def test_idor_other_bu_cannot_see_project(client_for):
    owner = client_for("E1001")
    pid = _demo_id(owner)
    other = client_for("E2001")  # 디스플레이사업부, 멤버 아님
    assert other.get(f"/api/projects/{pid}").status_code == 404
    assert other.get(f"/api/projects/{pid}/runs").status_code == 404
    assert other.post(f"/api/projects/{pid}/analysis", json={}).status_code == 404
    assert other.get(f"/api/projects/{pid}/runsheet.xlsx").status_code == 404


def test_role_matrix(client_for):
    owner = client_for("E1001")
    pid = _demo_id(owner)
    runner = client_for("E1002")
    viewer = client_for("E1003")
    # 열람자는 조회 가능, 결과 입력 불가
    assert viewer.get(f"/api/projects/{pid}/runs").status_code == 200
    run = viewer.get(f"/api/projects/{pid}/runs?status=planned").json()[0]
    r = viewer.post(f"/api/projects/{pid}/results", json={"rows": [{"run_id": run["id"], "values": {"etch_rate": 310}}]})
    assert r.status_code == 403
    # 실험자는 결과 입력 가능, 설정 변경·제안 불가
    r = runner.post(f"/api/projects/{pid}/results",
                    json={"rows": [{"run_id": run["id"], "values": {"etch_rate": 318.5, "uniformity": 2.4}}]})
    assert r.status_code == 200, r.text
    saved = r.json()["runs"][0]
    assert saved["status"] == "done" and saved["values"]["etch_rate"] == 318.5
    assert runner.patch(f"/api/projects/{pid}", json={"name": "x"}).status_code == 403
    assert runner.post(f"/api/projects/{pid}/recommend", json={}).status_code == 403
    # 공유 관리는 소유자만
    assert runner.get(f"/api/projects/{pid}/shares").status_code == 403


def test_measurement_versioning(client_for):
    owner = client_for("E1001")
    pid = _demo_id(owner)
    run = owner.get(f"/api/projects/{pid}/runs?status=done").json()[0]
    rid = run["id"]
    r = owner.post(f"/api/projects/{pid}/results", json={"rows": [{"run_id": rid, "values": {"etch_rate": 299.0}}]})
    assert r.status_code == 200
    r = owner.post(f"/api/projects/{pid}/results", json={"rows": [{"run_id": rid, "values": {"etch_rate": 301.0}}]})
    assert r.json()["runs"][0]["values"]["etch_rate"] == 301.0
    from app.db import SessionLocal
    from app.models import Measurement
    from sqlalchemy import select
    with SessionLocal() as db:
        ms = list(db.scalars(select(Measurement).where(Measurement.run_id == rid,
                                                       Measurement.response_key == "etch_rate")))
        assert len(ms) >= 3 and sum(m.is_current for m in ms) == 1


def test_analysis_recommend_accept_flow(client_for):
    owner = client_for("E1001")
    pid = _demo_id(owner)
    a = owner.post(f"/api/projects/{pid}/analysis", json={"validate_model": True})
    assert a.status_code == 200, a.text
    body = a.json()
    assert body["best"]["x"] and body["validation"]["available"]
    rec = owner.post(f"/api/projects/{pid}/recommend", json={"batch_size": 2})
    assert rec.status_code == 200, rec.text
    props = rec.json()["proposals"]
    assert len(props) == 2
    acc = owner.post(f"/api/projects/{pid}/batches/accept",
                     json={"proposals": props, "surrogate": "gp", "surrogate_version": rec.json()["surrogate_version"]})
    assert acc.status_code == 200 and acc.json()["runs_total"] == 2
    s = owner.post(f"/api/projects/{pid}/surface", json={"x_factor": "rf_power", "y_factor": "pressure",
                                                         "resolution": 12})
    assert s.status_code == 200 and len(s.json()["mean"]) == 12


def test_tabpfn_unavailable_is_reported_not_crash(client_for):
    owner = client_for("E1001")
    pid = _demo_id(owner)
    r = owner.post(f"/api/projects/{pid}/analysis", json={"surrogate": "tabpfn", "validate_model": False})
    assert r.status_code == 409
    cmp = owner.post(f"/api/projects/{pid}/compare", json={})
    assert cmp.status_code == 200
    models = {m["name"]: m for m in cmp.json()["models"]}
    assert models["gp"]["available"] and not models["tabpfn"]["available"]


def test_share_targeting_and_revoke(client_for):
    owner = client_for("E1001")
    pid = _demo_id(owner)
    depo = client_for("E1003")   # 증착공정개발팀 (시드에서 부서 공유됨)
    disp = client_for("E2001")   # 다른 사업부
    received = depo.get("/api/shares/received").json()
    assert any(s["project_id"] == pid for s in received)
    token = next(s["token"] for s in received if s["project_id"] == pid)
    v = depo.get(f"/api/shared/{token}")
    assert v.status_code == 200 and v.json()["best"]["x"]
    assert v.json()["raw"] is None  # 원본 데이터 숨김 기본값
    assert disp.get(f"/api/shared/{token}").status_code == 403
    assert depo.post(f"/api/shared/{token}/predict",
                     json={"points": [{"rf_power": 500, "pressure": 20, "cl2_flow": 80}]}).status_code == 200
    # 사업부 단위 공유 생성 후 철회
    from app.db import SessionLocal
    from app.models import BusinessUnit
    from sqlalchemy import select
    with SessionLocal() as db:
        disp_bu = db.scalar(select(BusinessUnit).where(BusinessUnit.code == "DISP")).id
    sh = owner.post(f"/api/projects/{pid}/shares", json={"target_type": "business_unit", "target_id": disp_bu,
                                                        "permission": "view", "mode": "snapshot"})
    assert sh.status_code == 200, sh.text
    t2 = sh.json()["token"]
    assert disp.get(f"/api/shared/{t2}").status_code == 200
    assert disp.post(f"/api/shared/{t2}/predict",
                     json={"points": [{"rf_power": 500, "pressure": 20, "cl2_flow": 80}]}).status_code == 403
    owner.delete(f"/api/projects/{pid}/shares/{sh.json()['id']}")
    assert disp.get(f"/api/shared/{t2}").status_code == 410


def test_soft_delete_and_restore(client_for):
    owner = client_for("E1001")
    r = owner.post("/api/projects", json={"name": "임시", "config": {
        "factors": [{"key": "a", "name": "A", "low": 0, "high": 10, "step": 1}],
        "responses": [{"key": "y", "name": "Y"}]}})
    assert r.status_code == 200, r.text
    pid = r.json()["id"]
    member = client_for("E1002")
    owner.post(f"/api/projects/{pid}/members", json={"user_id": member.get("/api/auth/me").json()["id"],
                                                     "role": "editor"})
    assert member.get(f"/api/projects/{pid}").status_code == 200
    owner.delete(f"/api/projects/{pid}")
    assert member.get(f"/api/projects/{pid}").status_code == 404
    assert any(p["id"] == pid for p in owner.get("/api/projects?scope=trash").json())
    assert owner.post(f"/api/projects/{pid}/restore").status_code == 200
    assert member.get(f"/api/projects/{pid}").status_code == 200


def test_mock_auth_refused_in_prod():
    with pytest.raises(RuntimeError):
        validate_settings(Settings(env="prod", auth_mode="mock", secret_key="x" * 32))


def test_excel_roundtrip_preview(client_for):
    owner = client_for("E1001")
    pid = _demo_id(owner)
    x = owner.get(f"/api/projects/{pid}/runsheet.xlsx")
    assert x.status_code == 200
    import io
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(x.content))
    ws = wb["실험 시트"]
    header = [c.value for c in ws[1]]
    col = header.index("식각률 [nm/min]") + 1
    for row in range(2, ws.max_row + 1):
        if ws.cell(row=row, column=header.index("상태") + 1).value == "계획":
            ws.cell(row=row, column=col).value = "3l5"  # 오타
            break
    buf = io.BytesIO()
    wb.save(buf)
    r = owner.post(f"/api/projects/{pid}/results/preview",
                   files={"file": ("sheet.xlsx", buf.getvalue(),
                                   "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert r.status_code == 200, r.text
    assert r.json()["summary"]["errors"] >= 1


def test_user_search_requires_query(client_for):
    c = client_for("E2001")
    assert c.get("/api/users/search?q=").json() == []
    assert c.get("/api/users/search?q=%20").json() == []
    assert [u["user_key"] for u in c.get("/api/users/search?q=김서연").json()] == ["E1001"]


def test_proposal_matching_pending_run_is_marked_replicate(client_for):
    """데모 초기 상태에서는 2차 배치(진행 중) 런과 같은 조건이 다시 제안된다. 이때 '신규'가 아니라 반복으로 표시해야 한다."""
    seed(reset=True)  # 다른 테스트가 만든 배치 영향 없이 데모 초기 상태로
    owner = client_for("E1001")
    pid = _demo_id(owner)
    pending = [r["planned"] for r in owner.get(f"/api/projects/{pid}/runs").json() if r["status"] == "planned"]
    props = owner.post(f"/api/projects/{pid}/recommend", json={"batch_size": 4}).json()["proposals"]
    dup = [p for p in props if p["x"] in pending]
    assert dup, "데모 데이터에서는 진행 중 런과 같은 조건이 제안되는 상황이 재현되어야 함"
    for p in dup:
        assert p["kind"] == "replicate" and "진행 중 런" in p["reason"]
    for p in props:  # 제안 예측값은 실제 데이터 기준이므로 산포가 0 근처로 붕괴하지 않음
        assert p["sigma"] > 3.0 and p["mean_hi"] - p["mean_lo"] > 5.0


def test_favorite_toggle_and_markers(client_for):
    owner = client_for("E1001")
    pid = _demo_id(owner)
    assert owner.put(f"/api/projects/{pid}/favorite").json() == {"favorite": True}
    owner.put(f"/api/projects/{pid}/favorite")  # 두 번 눌러도 하나만
    p = next(x for x in owner.get("/api/projects?scope=mine").json() if x["id"] == pid)
    assert p["is_favorite"] and p["member_count"] >= 1 and p["active_shares"] >= 1
    # 즐겨찾기는 개인 설정: 같은 프로젝트의 다른 멤버에게는 보이지 않음, 공유 수는 소유자에게만
    runner = client_for("E1002")
    q = next(x for x in runner.get("/api/projects?scope=member").json() if x["id"] == pid)
    assert not q["is_favorite"] and q["active_shares"] == 0
    assert owner.delete(f"/api/projects/{pid}/favorite").json() == {"favorite": False}
    assert not owner.get(f"/api/projects/{pid}").json()["is_favorite"]


def test_favorite_requires_access(client_for):
    owner = client_for("E1001")
    pid = _demo_id(owner)
    other = client_for("E2001")  # 다른 사업부, 멤버 아님
    assert other.put(f"/api/projects/{pid}/favorite").status_code == 404
    assert other.delete(f"/api/projects/{pid}/favorite").status_code == 404


def test_multi_objective_endpoints(client_for):
    seed(reset=True)
    owner = client_for("E1001")
    pid = _demo_id(owner)
    r = owner.post(f"/api/projects/{pid}/optimize", json={"validate_model": False})
    assert r.status_code == 200, r.text
    body = r.json()
    assert {d["key"] for d in body["responses"]} == {"etch_rate", "uniformity"}
    assert {x["key"] for x in body["best"]["responses"]} == {"etch_rate", "uniformity"}
    assert 0 <= body["best"]["desirability"] <= 1.5
    rec = owner.post(f"/api/projects/{pid}/recommend-multi", json={"batch_size": 3})
    assert rec.status_code == 200, rec.text
    assert len(rec.json()["proposals"]) == 3
    # 권한: 열람자는 결과는 보지만 제안은 못 받는다, 다른 사업부는 존재 자체를 모른다
    viewer = client_for("E1003")
    assert viewer.post(f"/api/projects/{pid}/optimize", json={"validate_model": False}).status_code == 200
    assert viewer.post(f"/api/projects/{pid}/recommend-multi", json={}).status_code == 403
    other = client_for("E2001")
    assert other.post(f"/api/projects/{pid}/optimize", json={}).status_code == 404


def test_add_factor_to_project_with_data(client_for):
    seed(reset=True)
    owner = client_for("E1001")
    pid = _demo_id(owner)
    p = owner.get(f"/api/projects/{pid}").json()
    cfg = p["config"]
    cfg["factors"].append({"key": "temp", "name": "척 온도", "unit": "°C", "low": 20, "high": 80, "step": 1, "scale": "linear"})
    # 기존 실험 값이 없으면 거부
    r = owner.patch(f"/api/projects/{pid}", json={"config": cfg})
    assert r.status_code == 422 and "기존 실험 값" in r.json()["detail"]
    # 범위 밖이면 거부
    assert owner.patch(f"/api/projects/{pid}", json={"config": cfg, "new_factor_values": {"temp": 200}}).status_code == 422
    before = owner.post(f"/api/projects/{pid}/optimize", json={"validate_model": False}).json()["responses"][0]["data"]["n_obs"]
    r = owner.patch(f"/api/projects/{pid}", json={"config": cfg, "new_factor_values": {"temp": 40}})
    assert r.status_code == 200, r.text
    runs = owner.get(f"/api/projects/{pid}/runs").json()
    assert all(run["planned"]["temp"] == 40 for run in runs)
    # 기존 데이터가 학습에서 빠지지 않는다
    after = owner.post(f"/api/projects/{pid}/optimize", json={"validate_model": False}).json()
    assert after["responses"][0]["data"]["n_obs"] == before
    assert "temp" in after["best"]["x"]
    seed(reset=True)


def test_second_demo_multi_response(client_for):
    seed(reset=True)
    owner = client_for("E1001")
    ps = owner.get("/api/projects?scope=mine").json()
    cmp = next(p for p in ps if p["name"] == "CMP 슬러리 배합 최적화 (데모)")
    assert cmp["runs_done"] > 0 and cmp["runs_open"] == 0  # 첫 DOE 결과가 모두 입력된 상태 → 바로 ② 학습 결과
    body = owner.post(f"/api/projects/{cmp['id']}/optimize", json={"validate_model": False}).json()
    assert len(body["best"]["responses"]) == 3
    assert body["best"]["desirability"] < 0.95, "응답끼리 충돌하므로 모두 완벽히 만족하지는 못해야 예제로서 의미가 있음"
    assert {a["why"] for a in body["alternatives"]} >= {"제거율 우선", "디싱 우선", "결함 수 우선"}


def test_predict_multi_all_responses(client_for):
    seed(reset=True)
    owner = client_for("E1001")
    pid = _demo_id(owner)
    x = {"rf_power": 530, "pressure": 41, "cl2_flow": 75}
    r = owner.post(f"/api/projects/{pid}/predict-multi", json={"points": [x, {**x, "pressure": 79}]})
    assert r.status_code == 200, r.text
    a, b = r.json()
    assert {p["key"] for p in a["responses"]} == {"etch_rate", "uniformity"}
    assert 0 <= a["desirability"] <= 1.5
    assert a["desirability"] > b["desirability"], "압력이 높으면 산포가 커져 목표 달성이 떨어져야 함(데모 공정)"
    other = client_for("E2001")
    assert other.post(f"/api/projects/{pid}/predict-multi", json={"points": [x]}).status_code == 404


def test_run_codes_are_korean_round_format(client_for):
    seed(reset=True)
    owner = client_for("E1001")
    pid = _demo_id(owner)
    codes = [r["code"] for r in owner.get(f"/api/projects/{pid}/runs").json()]
    assert "1차-01" in codes and "2차-01" in codes and not any(c.startswith("B") for c in codes)
    # 예전 형식이 남아 있던 DB도 시작할 때 바뀐다
    from app.db import SessionLocal
    from app.migrations import migrate_run_codes
    from app.models import Run
    with SessionLocal() as db:
        run = db.query(Run).filter(Run.code == "1차-01").first()
        run.code = "B1-01"
        db.commit()
        assert migrate_run_codes(db) == 1
        assert db.get(Run, run.id).code == "1차-01"


def test_clearing_results_returns_run_to_planned(client_for):
    seed(reset=True)
    owner = client_for("E1001")
    pid = _demo_id(owner)
    run = next(r for r in owner.get(f"/api/projects/{pid}/runs").json() if r["status"] == "planned")
    r = owner.post(f"/api/projects/{pid}/results", json={"rows": [{"run_id": run["id"], "values": {"etch_rate": 320, "uniformity": 3}}]})
    assert next(x for x in r.json()["runs"])["status"] == "done"
    r = owner.post(f"/api/projects/{pid}/results", json={"rows": [{"run_id": run["id"], "values": {"etch_rate": None}}]})
    assert next(x for x in r.json()["runs"])["status"] == "planned"


def test_delete_run_soft_and_authorized(client_for):
    seed(reset=True)
    owner = client_for("E1001")
    pid = _demo_id(owner)
    runs = owner.get(f"/api/projects/{pid}/runs").json()
    done = next(r for r in runs if r["status"] == "done")
    before = owner.get(f"/api/projects/{pid}").json()
    runner = client_for("E1002")
    assert runner.delete(f"/api/projects/{pid}/runs/{done['id']}").status_code == 403  # 실험자는 삭제 불가
    other = client_for("E2001")
    assert other.delete(f"/api/projects/{pid}/runs/{done['id']}").status_code == 404
    assert owner.delete(f"/api/projects/{pid}/runs/{done['id']}").status_code == 200
    after_runs = owner.get(f"/api/projects/{pid}/runs").json()
    assert done["id"] not in {r["id"] for r in after_runs}
    after = owner.get(f"/api/projects/{pid}").json()
    assert after["runs_total"] == before["runs_total"] - 1 and after["runs_done"] == before["runs_done"] - 1
    assert owner.delete(f"/api/projects/{pid}/runs/{done['id']}").status_code == 404  # 두 번 삭제 불가
    # 기록은 남는다 (복구 가능)
    from app.db import SessionLocal
    from app.models import AuditLog, Measurement, Run
    with SessionLocal() as db:
        assert db.get(Run, done["id"]).status == "excluded"
        assert db.query(Measurement).filter(Measurement.run_id == done["id"]).count() > 0
        assert db.query(AuditLog).filter(AuditLog.action == "run.delete").count() == 1


def _plain_project(c, name="가져오기 테스트"):
    r = c.post("/api/projects", json={"name": name, "config": {
        "factors": [{"key": "temp", "name": "온도", "unit": "°C", "low": 100, "high": 200, "step": 1},
                    {"key": "time", "name": "시간", "unit": "min", "low": 10, "high": 60, "step": 5}],
        "responses": [{"key": "yield", "name": "수율", "unit": "%", "goal": "maximize"}],
        "settings": {"batch_size": 3, "budget_runs": 30}}})
    assert r.status_code == 200, r.text
    return r.json()["id"]


def test_import_existing_data(client_for):
    owner = client_for("E1001")
    pid = _plain_project(owner)
    rows = [{"x": {"temp": 120, "time": 20}, "values": {"yield": 61.2}, "note": "작년 3월"},
            {"x": {"temp": 120, "time": 20}, "values": {"yield": 63.0}},           # 같은 조건 → 반복
            {"x": {"temp": 173.28, "time": 33}, "values": {"yield": 70.4}},       # 세팅 정밀도와 달라도 그대로
            {"x": {"temp": 190, "time": 55}, "values": {}}]                        # 결과 없음 → 대기
    r = owner.post(f"/api/projects/{pid}/import", json={"rows": rows})
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["imported"] == 4 and body["done"] == 3 and body["replicated_conditions"] == 1
    assert body["batch"]["seq"] == 0
    runs = owner.get(f"/api/projects/{pid}/runs").json()
    assert [x["code"] for x in runs] == ["0차-01", "0차-02", "0차-03", "0차-04"]
    assert [x["replicate_no"] for x in runs] == [1, 2, 1, 1]
    assert runs[2]["actual"]["temp"] == 173.28
    assert runs[0]["deviation_note"] == "작년 3월"
    assert [x["status"] for x in runs] == ["done", "done", "done", "planned"]
    # 이어서 더 가져오면 같은 0차에 번호가 이어지고, 앞서 가져온 같은 조건과 반복으로 묶인다
    r = owner.post(f"/api/projects/{pid}/import", json={"rows": [{"x": {"temp": 120, "time": 20}, "values": {"yield": 62}}]})
    assert r.status_code == 200
    runs = owner.get(f"/api/projects/{pid}/runs").json()
    assert runs[-1]["code"] == "0차-05" and runs[-1]["replicate_no"] == 3
    # 바로 학습할 수 있고, 다음 제안은 1차부터
    opt = owner.post(f"/api/projects/{pid}/optimize", json={"validate_model": False})
    assert opt.status_code == 200, opt.text
    assert opt.json()["responses"][0]["data"]["n_obs"] == 4  # 결과가 있는 런만 (0차-04는 결과 없음)
    rec = owner.post(f"/api/projects/{pid}/recommend-multi", json={"batch_size": 2})
    assert rec.status_code == 200, rec.text
    acc = owner.post(f"/api/projects/{pid}/batches/manual", json={"points": [{"temp": 150, "time": 30}]})
    assert acc.json()["seq"] == 1


def test_import_rejects_bad_rows_and_needs_editor(client_for):
    owner = client_for("E1001")
    pid = _plain_project(owner, "가져오기 거부")
    r = owner.post(f"/api/projects/{pid}/import", json={"rows": [{"x": {"temp": 250, "time": 20}, "values": {"yield": 1}}]})
    assert r.status_code == 422 and "범위" in r.json()["detail"]
    r = owner.post(f"/api/projects/{pid}/import", json={"rows": [{"x": {"temp": 150}, "values": {}}]})
    assert r.status_code == 422 and "시간" in r.json()["detail"]
    r = owner.post(f"/api/projects/{pid}/import", json={"rows": [{"x": {"temp": 150, "time": 20}, "values": {"zzz": 1}}]})
    assert r.status_code == 422
    assert owner.get(f"/api/projects/{pid}/runs").json() == []  # 하나라도 틀리면 아무것도 들어가지 않음
    # 다른 사업부 사용자는 존재도 모름(404), 열람자는 403
    assert client_for("E2001").post(f"/api/projects/{pid}/import", json={"rows": [{"x": {"temp": 150, "time": 20}}]}).status_code == 404
    viewer = client_for("E1003")
    owner.post(f"/api/projects/{pid}/members", json={"user_id": viewer.get("/api/auth/me").json()["id"], "role": "viewer"})
    assert viewer.post(f"/api/projects/{pid}/import", json={"rows": [{"x": {"temp": 150, "time": 20}}]}).status_code == 403


def test_export_runs_xlsx(client_for):
    import io as _io
    from openpyxl import load_workbook
    owner = client_for("E1001")
    pid = _plain_project(owner, "엑셀 다운로드")
    owner.post(f"/api/projects/{pid}/import", json={"rows": [
        {"x": {"temp": 120, "time": 20}, "values": {"yield": 61.25}, "note": "메모1"},
        {"x": {"temp": 120, "time": 20}, "values": {"yield": 63}},
        {"x": {"temp": 150, "time": 30}, "values": {}}]})
    rid = owner.get(f"/api/projects/{pid}/runs").json()[2]["id"]
    owner.delete(f"/api/projects/{pid}/runs/{rid}")  # 삭제한 런은 빠짐
    r = owner.get(f"/api/projects/{pid}/runs.xlsx")
    assert r.status_code == 200
    assert "filename*=UTF-8''" in r.headers["content-disposition"]
    wb = load_workbook(_io.BytesIO(r.content))
    rows = list(wb["실험 데이터"].iter_rows(values_only=True))
    assert rows[0] == ("런 ID", "순서", "온도 [°C]", "시간 [min]", "수율 [%]", "반복", "메모", "상태")
    assert rows[1] == ("0차-01", 1, 120, 20, 61.25, "1/2", "메모1", "완료")
    assert len(rows) == 3
    assert wb["DOE 설정"]["B1"].value == "엑셀 다운로드"
    # 접근 권한 없는 사람은 존재도 모름
    assert client_for("E2001").get(f"/api/projects/{pid}/runs.xlsx").status_code == 404
