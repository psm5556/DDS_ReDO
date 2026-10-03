"""실험 시트 엑셀 왕복 (CLAUDE.md 7.2절): 잠긴 ID 열 + 데이터 유효성, 업로드 미리보기."""
from __future__ import annotations

import csv
import io
import math

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from ..models import Measurement, Project, Run
from ..schemas import ProjectConfig
from .results import STATUS_FROM_KO, STATUS_KO

LOCK_FILL = PatternFill("solid", fgColor="E8ECF1")
EDIT_FILL = PatternFill("solid", fgColor="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")


def _headers(cfg: ProjectConfig) -> list[tuple[str, str, bool]]:
    """(헤더 텍스트, 내부 키, 편집 가능)"""
    h: list[tuple[str, str, bool]] = [("런 ID", "code", False), ("실행 순서", "order", False), ("배치", "batch", False),
                                      ("반복", "rep", False)]
    for f in cfg.factors:
        h.append((f"{f.name} 계획 [{f.unit}]" if f.unit else f"{f.name} 계획", f"plan:{f.key}", False))
    for f in cfg.factors:
        h.append((f"{f.name} 실제 [{f.unit}]" if f.unit else f"{f.name} 실제", f"act:{f.key}", True))
    for r in cfg.responses:
        h.append((f"{r.name} [{r.unit}]" if r.unit else r.name, f"val:{r.key}", True))
    h += [("상태", "status", True), ("메모", "note", True)]
    return h


def build_runsheet(project: Project, cfg: ProjectConfig, runs: list[Run], batch_seq: dict[int, int],
                   meas: dict[tuple[int, str], Measurement]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "실험 시트"
    heads = _headers(cfg)
    ws.append([h for h, _, _ in heads])
    for c, (_, _, editable) in enumerate(heads, start=1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(c)].width = 16
    for run in runs:
        row: list = []
        for _, key, _ in heads:
            if key == "code":
                row.append(run.code)
            elif key == "order":
                row.append(run.run_order)
            elif key == "batch":
                row.append(batch_seq.get(run.batch_id))
            elif key == "rep":
                row.append(run.replicate_no)
            elif key.startswith("plan:"):
                row.append(run.planned.get(key[5:]))
            elif key.startswith("act:"):
                row.append((run.actual or run.planned).get(key[4:]))
            elif key.startswith("val:"):
                m = meas.get((run.id, key[4:]))
                row.append(m.value if m else None)
            elif key == "status":
                row.append(STATUS_KO.get(run.status, run.status))
            elif key == "note":
                row.append(run.deviation_note or "")
        ws.append(row)
    n = len(runs) + 1
    for c, (_, key, editable) in enumerate(heads, start=1):
        for r in range(2, n + 1):
            cell = ws.cell(row=r, column=c)
            cell.protection = Protection(locked=not editable)
            cell.fill = EDIT_FILL if editable else LOCK_FILL
        col = get_column_letter(c)
        if key == "status":
            dv = DataValidation(type="list", formula1='"계획,진행중,완료,실패,실행불가"', allow_blank=False)
            dv.add(f"{col}2:{col}{max(n, 2)}")
            ws.add_data_validation(dv)
        elif key.startswith("val:") or key.startswith("act:"):
            dv = DataValidation(type="decimal", operator="between", formula1="-1E+307", formula2="1E+307",
                                allow_blank=True, error="숫자만 입력할 수 있습니다.", errorTitle="입력 오류")
            dv.add(f"{col}2:{col}{max(n, 2)}")
            ws.add_data_validation(dv)
    ws.freeze_panes = "B2"
    ws.protection.sheet = True
    ws.protection.formatColumns = False
    ws.protection.autoFilter = False
    meta = wb.create_sheet("_meta")
    meta.append(["project_id", project.id])
    meta.append(["keys"] + [k for _, k, _ in heads])
    meta.sheet_state = "hidden"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _num(v: object) -> float | None | str:
    if v is None or (isinstance(v, str) and v.strip() == ""):
        return None
    if isinstance(v, (int, float)):
        return float(v) if math.isfinite(float(v)) else "invalid"
    try:
        return float(str(v).replace(",", "").strip())
    except ValueError:
        return "invalid"


def parse_upload(filename: str, content: bytes, project: Project, cfg: ProjectConfig) -> tuple[list[dict], list[str]]:
    """업로드 파일 → [{code, actual, values, status, note, errors}] , 파일 수준 오류."""
    heads = _headers(cfg)
    label_to_key = {h: k for h, k, _ in heads}
    rows: list[list] = []
    file_errors: list[str] = []
    if filename.lower().endswith((".xlsx", ".xlsm")):
        wb = load_workbook(io.BytesIO(content), data_only=True)
        if "_meta" in wb.sheetnames:
            pid = wb["_meta"].cell(row=1, column=2).value
            if pid is not None and int(pid) != project.id:
                file_errors.append("다른 프로젝트의 실험 시트입니다. 이 프로젝트에서 내려받은 시트를 사용하세요.")
        ws = wb["실험 시트"] if "실험 시트" in wb.sheetnames else wb.active
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
    elif filename.lower().endswith(".csv"):
        text = content.decode("utf-8-sig", errors="replace")
        rows = [list(r) for r in csv.reader(io.StringIO(text))]
    else:
        return [], ["엑셀(.xlsx) 또는 CSV 파일만 올릴 수 있습니다."]
    if not rows:
        return [], ["파일이 비어 있습니다."]
    header = [str(h).strip() if h is not None else "" for h in rows[0]]
    keys = [label_to_key.get(h) for h in header]
    if "code" not in keys:
        return [], ["'런 ID' 열을 찾을 수 없습니다. 앱에서 내려받은 실험 시트 양식을 사용하세요."]
    out: list[dict] = []
    for raw in rows[1:]:
        if all(v is None or str(v).strip() == "" for v in raw):
            continue
        rec: dict = {"code": None, "actual": {}, "values": {}, "status": None, "note": None, "errors": []}
        for k, v in zip(keys, raw):
            if k is None:
                continue
            if k == "code":
                rec["code"] = str(v).strip() if v is not None else None
            elif k.startswith("act:"):
                n = _num(v)
                if n == "invalid":
                    rec["errors"].append(f"{k[4:]} 실제값이 숫자가 아닙니다: {v}")
                elif n is not None:
                    rec["actual"][k[4:]] = n
            elif k.startswith("val:"):
                n = _num(v)
                if n == "invalid":
                    rec["errors"].append(f"{k[4:]} 결과값이 숫자가 아닙니다: {v}")
                else:
                    rec["values"][k[4:]] = n
            elif k == "status" and v is not None and str(v).strip():
                s = STATUS_FROM_KO.get(str(v).strip(), str(v).strip())
                if s not in ("planned", "running", "done", "failed", "infeasible"):
                    rec["errors"].append(f"상태 값을 알 수 없습니다: {v}")
                else:
                    rec["status"] = s
            elif k == "note":
                rec["note"] = "" if v is None else str(v)
        out.append(rec)
    return out, file_errors
